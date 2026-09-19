from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


# =========================================================
# Project paths
# =========================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

TABLE_DIR = PROJECT_ROOT / "output" / "tables"
FIGURE_DIR = PROJECT_ROOT / "output" / "figures"
REPORT_DIR = PROJECT_ROOT / "report"
ARTIFACT_DIR = PROJECT_ROOT / "validation" / "artifacts"

STATISTICS_FILE = (
    TABLE_DIR
    / "bkm_statistics_all_samples.csv"
)

INTERPRETATION_FILE = (
    REPORT_DIR
    / "economic_interpretation.md"
)

QA_FILE = (
    ARTIFACT_DIR
    / "quality_assurance.json"
)

INTERPRETATION_ARTIFACT = (
    ARTIFACT_DIR
    / "economic_interpretation.json"
)

MARKDOWN_OUTPUT = (
    REPORT_DIR
    / "final_report.md"
)

DOCX_OUTPUT = (
    REPORT_DIR
    / "final_report.docx"
)

FINAL_ARTIFACT = (
    ARTIFACT_DIR
    / "final_report.json"
)


FIGURE_FILES = [
    FIGURE_DIR / "time_series_SP500.png",
    FIGURE_DIR / "time_series_Small_Growth.png",
    FIGURE_DIR / "time_series_Small_Value.png",
    FIGURE_DIR / "time_series_Big_Growth.png",
    FIGURE_DIR / "time_series_Big_Value.png",
    FIGURE_DIR / "histogram_normal_SP500.png",
    FIGURE_DIR / "histogram_normal_Small_Growth.png",
    FIGURE_DIR / "histogram_normal_Small_Value.png",
    FIGURE_DIR / "histogram_normal_Big_Growth.png",
    FIGURE_DIR / "histogram_normal_Big_Value.png",
    FIGURE_DIR / "tail_comparison_ecdf.png",
]


# =========================================================
# Utilities
# =========================================================

def ensure_directories() -> None:

    REPORT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    ARTIFACT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )


def sha256_file(
    path: Path,
) -> str:

    digest = hashlib.sha256()

    with path.open(
        "rb"
    ) as file:

        for chunk in iter(
            lambda: file.read(65536),
            b"",
        ):

            digest.update(
                chunk
            )

    return digest.hexdigest()


def read_json(
    path: Path,
) -> dict:

    if not path.is_file():

        raise FileNotFoundError(
            f"Required file missing: "
            f"{path.relative_to(PROJECT_ROOT)}"
        )

    with path.open(
        "r",
        encoding="utf-8",
    ) as file:

        return json.load(file)


def read_text(
    path: Path,
) -> str:

    if not path.is_file():

        raise FileNotFoundError(
            f"Required file missing: "
            f"{path.relative_to(PROJECT_ROOT)}"
        )

    return path.read_text(
        encoding="utf-8"
    )


def read_statistics() -> pd.DataFrame:

    if not STATISTICS_FILE.is_file():

        raise FileNotFoundError(
            "Missing statistical results table."
        )

    return pd.read_csv(
        STATISTICS_FILE
    )


# =========================================================
# Preconditions
# =========================================================

def require_upstream_passes() -> None:

    qa = read_json(
        QA_FILE
    )

    interpretation = read_json(
        INTERPRETATION_ARTIFACT
    )

    if qa.get("status") != "passed":

        raise RuntimeError(
            "Quality assurance has not passed."
        )

    if interpretation.get("status") != "passed":

        raise RuntimeError(
            "Economic interpretation has not passed."
        )


def require_figures() -> None:

    missing = [
        str(
            path.relative_to(
                PROJECT_ROOT
            )
        )
        for path in FIGURE_FILES
        if not path.is_file()
    ]

    if missing:

        raise FileNotFoundError(
            f"Required figures missing: {missing}"
        )


# =========================================================
# Formatting helpers
# =========================================================

def format_percent(
    value: float,
    decimals: int = 2,
) -> str:

    return (
        f"{value * 100:.{decimals}f}%"
    )


def format_number(
    value: float,
    decimals: int = 3,
) -> str:

    return (
        f"{value:.{decimals}f}"
    )


def format_statistics_table(
    df: pd.DataFrame,
) -> pd.DataFrame:

    formatted = df.copy()

    percent_columns = [
        "Mean_Excess_Annualized",
        "Std_Annualized",
        "Lower_Partial_Std_Annualized",
        "Empirical_VaR_1pct_Monthly",
        "Normal_VaR_1pct_Monthly",
        "Expected_Shortfall_1pct_Monthly",
    ]

    number_columns = [
        "Sharpe_Ratio",
        "Skewness",
        "Excess_Kurtosis",
    ]

    for column in percent_columns:

        formatted[column] = (
            formatted[column]
            .apply(
                lambda x: format_percent(
                    float(x)
                )
            )
        )

    for column in number_columns:

        formatted[column] = (
            formatted[column]
            .apply(
                lambda x: format_number(
                    float(x)
                )
            )
        )

    formatted[
        "Pct_Below_Mean_Minus_3SD"
    ] = (
        formatted[
            "Pct_Below_Mean_Minus_3SD"
        ]
        .apply(
            lambda x: (
                f"{float(x):.2f}%"
            )
        )
    )

    return formatted


# =========================================================
# Markdown report
# =========================================================

def dataframe_to_markdown(
    df: pd.DataFrame,
) -> str:

    return df.to_markdown(
        index=False
    )


def build_markdown_report(
    statistics: pd.DataFrame,
    interpretation: str,
) -> str:

    formatted = format_statistics_table(
        statistics
    )

    sections = []

    sections.append(
        "# BUFN741 Group Assignment 1\n"
    )

    sections.append(
        "## Return Distributions, Tail Risk, "
        "and Equity Style Portfolios\n"
    )

    sections.append(
        "## 1. Research Design\n\n"
        "The analysis uses monthly returns for the S&P 500 "
        "and four Fama-French 2×3 Size/Book-to-Market "
        "portfolios: Small/Growth, Small/Value, "
        "Big/Growth, and Big/Value.\n\n"
        "Monthly simple excess returns are constructed as "
        "asset return minus the one-month risk-free rate. "
        "The analysis covers the full common sample and "
        "three subperiods: 1990–2000, 2001–2010, and "
        "2011–present.\n"
    )

    sections.append(
        "## 2. Statistical Results\n"
    )

    for sample in formatted[
        "Sample"
    ].unique():

        sample_df = (
            formatted[
                formatted["Sample"]
                == sample
            ]
            .copy()
        )

        sections.append(
            f"### {sample}\n"
        )

        sections.append(
            dataframe_to_markdown(
                sample_df
            )
        )

        sections.append(
            "\n"
        )

    sections.append(
        "## 3. Figures\n\n"
        "The complete figure set is stored in "
        "`output/figures/`. It includes monthly "
        "excess-return time series, empirical return "
        "histograms with fitted normal distributions, "
        "and a common left-tail ECDF comparison.\n"
    )

    sections.append(
        "## 4. Economic Interpretation\n\n"
    )

    sections.append(
        interpretation
    )

    sections.append(
        "\n\n## 5. Reproducibility and Quality Assurance\n\n"
        "The final report was assembled only after the "
        "data-acquisition, validation, processing, "
        "statistical-analysis, figure-generation, "
        "quality-assurance, and economic-interpretation "
        "stages completed successfully.\n\n"
        "Numerical statistics were generated by "
        "deterministic Python code. Economic interpretation "
        "was produced by Codex after the upstream quality "
        "assurance gate passed.\n"
    )

    return "\n".join(
        sections
    )


def write_markdown(
    text: str,
) -> None:

    MARKDOWN_OUTPUT.write_text(
        text + "\n",
        encoding="utf-8",
    )


# =========================================================
# Word report
# =========================================================

def add_title_page(
    document: Document,
) -> None:

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = title.add_run(
        "BUFN741 Group Assignment 1"
    )

    run.bold = True
    run.font.size = Pt(20)

    subtitle = document.add_paragraph()

    subtitle.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = subtitle.add_run(
        "Return Distributions, Tail Risk, "
        "and Equity Style Portfolios"
    )

    run.font.size = Pt(14)

    document.add_page_break()


def add_statistics_table(
    document: Document,
    df: pd.DataFrame,
) -> None:

    display_columns = [
        "Asset",
        "Usable_Months",
        "Mean_Excess_Annualized",
        "Std_Annualized",
        "Sharpe_Ratio",
        "Lower_Partial_Std_Annualized",
        "Skewness",
        "Excess_Kurtosis",
        "Empirical_VaR_1pct_Monthly",
        "Normal_VaR_1pct_Monthly",
        "Pct_Below_Mean_Minus_3SD",
        "Expected_Shortfall_1pct_Monthly",
    ]

    table = document.add_table(
        rows=1,
        cols=len(
            display_columns
        ),
    )

    table.style = (
        "Table Grid"
    )

    header = table.rows[0].cells

    short_headers = [
        "Asset",
        "N",
        "Ann. Mean",
        "Ann. SD",
        "Sharpe",
        "Downside SD",
        "Skew",
        "Excess Kurt.",
        "Emp. 1% VaR",
        "Normal 1% VaR",
        "% < Mean−3SD",
        "1% ES",
    ]

    for i, name in enumerate(
        short_headers
    ):

        header[i].text = name

    for _, row in df.iterrows():

        cells = table.add_row().cells

        for i, column in enumerate(
            display_columns
        ):

            cells[i].text = str(
                row[column]
            )


def add_figures(
    document: Document,
) -> None:

    document.add_heading(
        "Figures",
        level=1,
    )

    figure_groups = [
        (
            "Monthly Excess Returns",
            FIGURE_FILES[:5],
        ),
        (
            "Empirical versus Normal Distributions",
            FIGURE_FILES[5:10],
        ),
        (
            "Left-Tail Comparison",
            FIGURE_FILES[10:],
        ),
    ]

    for heading, files in (
        figure_groups
    ):

        document.add_heading(
            heading,
            level=2,
        )

        for path in files:

            document.add_picture(
                str(path),
                width=Inches(6.2),
            )

            caption = (
                document.add_paragraph()
            )

            caption.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            caption.add_run(
                path.stem
            ).italic = True


def add_interpretation(
    document: Document,
    interpretation: str,
) -> None:

    document.add_heading(
        "Economic Interpretation",
        level=1,
    )

    for line in interpretation.splitlines():

        stripped = line.strip()

        if not stripped:

            continue

        if stripped.startswith(
            "## "
        ):

            document.add_heading(
                stripped[3:],
                level=2,
            )

        elif stripped.startswith(
            "### "
        ):

            document.add_heading(
                stripped[4:],
                level=3,
            )

        else:

            document.add_paragraph(
                stripped
            )


def build_docx_report(
    statistics: pd.DataFrame,
    interpretation: str,
) -> None:

    document = Document()

    add_title_page(
        document
    )

    document.add_heading(
        "Research Design",
        level=1,
    )

    document.add_paragraph(
        "The analysis examines monthly simple excess "
        "returns for the S&P 500 and four Fama-French "
        "2×3 Size/Book-to-Market portfolios: "
        "Small/Growth, Small/Value, Big/Growth, "
        "and Big/Value."
    )

    document.add_paragraph(
        "The analysis is performed over the full common "
        "sample and the subperiods 1990–2000, "
        "2001–2010, and 2011–present."
    )

    document.add_heading(
        "Statistical Results",
        level=1,
    )

    formatted = format_statistics_table(
        statistics
    )

    for sample in formatted[
        "Sample"
    ].unique():

        document.add_heading(
            str(sample),
            level=2,
        )

        sample_df = (
            formatted[
                formatted["Sample"]
                == sample
            ]
        )

        add_statistics_table(
            document,
            sample_df,
        )

        document.add_paragraph()

    add_figures(
        document
    )

    add_interpretation(
        document,
        interpretation,
    )

    document.add_heading(
        "Reproducibility and Quality Assurance",
        level=1,
    )

    document.add_paragraph(
        "All numerical outputs were generated using "
        "deterministic Python code. The cross-stage "
        "quality-assurance stage passed before economic "
        "interpretation was generated."
    )

    document.add_paragraph(
        "Codex was used only for economic reasoning and "
        "interpretation after validated quantitative "
        "evidence had been produced."
    )

    document.save(
        DOCX_OUTPUT
    )


# =========================================================
# Final artifact
# =========================================================

def build_artifact(
    statistics: pd.DataFrame,
) -> dict:

    errors = []
    warnings = []

    if not MARKDOWN_OUTPUT.is_file():

        errors.append(
            "final_report.md was not created."
        )

    if not DOCX_OUTPUT.is_file():

        errors.append(
            "final_report.docx was not created."
        )

    if (
        MARKDOWN_OUTPUT.is_file()
        and MARKDOWN_OUTPUT.stat().st_size
        == 0
    ):

        errors.append(
            "final_report.md is empty."
        )

    if (
        DOCX_OUTPUT.is_file()
        and DOCX_OUTPUT.stat().st_size
        == 0
    ):

        errors.append(
            "final_report.docx is empty."
        )

    if len(statistics) != 20:

        errors.append(
            "Expected 20 statistical result rows."
        )

    status = (
        "passed"
        if not errors
        else "failed"
    )

    evidence_files = [
        STATISTICS_FILE,
        INTERPRETATION_FILE,
        QA_FILE,
        INTERPRETATION_ARTIFACT,
    ]

    evidence = []

    for path in evidence_files:

        evidence.append(
            {
                "file": str(
                    path.relative_to(
                        PROJECT_ROOT
                    )
                ),
                "sha256": sha256_file(
                    path
                ),
            }
        )

    return {
        "stage": "final_report",
        "status": status,
        "artifact_version": "1.0",
        "outputs": {
            "markdown": (
                "report/final_report.md"
            ),
            "word": (
                "report/final_report.docx"
            ),
        },
        "statistics_rows": int(
            len(statistics)
        ),
        "figure_count": len(
            FIGURE_FILES
        ),
        "evidence": evidence,
        "errors": errors,
        "warnings": warnings,
    }


def write_artifact(
    artifact: dict,
) -> None:

    FINAL_ARTIFACT.write_text(
        json.dumps(
            artifact,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )


# =========================================================
# Main
# =========================================================

def main() -> int:

    ensure_directories()

    try:

        require_upstream_passes()

        require_figures()

        statistics = (
            read_statistics()
        )

        interpretation = (
            read_text(
                INTERPRETATION_FILE
            )
        )

        markdown = (
            build_markdown_report(
                statistics,
                interpretation,
            )
        )

        write_markdown(
            markdown
        )

        build_docx_report(
            statistics,
            interpretation,
        )

        artifact = (
            build_artifact(
                statistics
            )
        )

        write_artifact(
            artifact
        )

        print()
        print("=" * 72)
        print("BUFN741 FINAL REPORT")
        print("=" * 72)

        print()
        print(
            f"Status: "
            f"{artifact['status'].upper()}"
        )

        print()
        print(
            "Outputs:"
        )

        print(
            "  report/final_report.md"
        )

        print(
            "  report/final_report.docx"
        )

        print()
        print(
            "Evidence artifact:"
        )

        print(
            "  validation/artifacts/"
            "final_report.json"
        )

        print("=" * 72)

        return (
            0
            if artifact["status"]
            == "passed"
            else 1
        )

    except Exception as exc:

        print()
        print("=" * 72)
        print("FINAL REPORT BUILD FAILED")
        print("=" * 72)

        print()
        print(
            type(exc).__name__
        )

        print(
            str(exc)
        )

        print()

        return 1


if __name__ == "__main__":
    sys.exit(main())